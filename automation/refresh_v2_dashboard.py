#!/usr/bin/env python3
"""
Refreshes progress_dashboard_v2.html (the Executive Program Board) from the
live "Showroom Revamp Roadmap.xlsx" workbook via Microsoft Graph (app-only).

Reads two sheets — "Showroom Plan V2" and "Old-New Sequence" — and rewrites the
single `const DATA = {...};` object the page consumes. Unlike index.html, that
payload is JSON-encoded, so multi-line cell values (excludedDates, agreedTeams)
survive verbatim and cannot corrupt the surrounding JavaScript.

Credentials come from the environment, same three repository secrets as the
index.html refresh:

    AZURE_TENANT_ID
    AZURE_CLIENT_ID
    AZURE_CLIENT_SECRET   (omit once a federated credential is configured)

Usage: refresh_v2_dashboard.py <in_html> <out_html>
"""
import datetime as dt
import json
import os
import re
import sys
from io import BytesIO

import openpyxl
import requests

# Drive item for "Showroom Revamp Roadmap.xlsx" (owner: mina_faham_com).
# Not secrets — Graph object IDs only resolve with an authorized token.
DRIVE_ID = "b!fH2okKMc606AhlGPvxRbqQIGlG9GCfNCn8EedNT7pWp6O0pSKU0HRo7S4etT9Wdk"
ITEM_ID = "01QNMY3327RKP43ZJ2IFFJR5QKBK64W47X"

PLAN_SHEET = 'Showroom Plan V2'
SEQ_SHEET = 'Old-New Sequence'
SOURCE_LABEL = ('Showroom Revamp Roadmap.xlsx — "Showroom Plan V2" '
                '+ "Old-New Sequence"')
EXPECTED_SHOWROOMS = 5

# "2g", "1g", "1T, 1g" — a team spec. Column D also carries free-text notes
# ("This room is not mentioned from the Design sheet"), so it must be told apart.
TEAM_RE = re.compile(r'^\s*\d+\s*[A-Za-z]{1,2}\s*(,\s*\d+\s*[A-Za-z]{1,2}\s*)*$')


# Column positions are NOT assumed. The workbook's columns have been inserted
# into and reordered in the past, which silently shifted every field by one and
# published wrong numbers while all the row-level gates still passed. Columns
# are therefore resolved by matching the header row, and a missing REQUIRED
# header aborts the run.
REQUIRED = ('showroom', 'phase', 'rooms',
            'plannedRemoval', 'plannedInstall', 'plannedCompletion')


def header_map(ws):
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    hdr = [' '.join(str(c).split()).lower() if c is not None else '' for c in row]

    def find(*subs, after=-1):
        for i, h in enumerate(hdr):
            if i <= after or not h:
                continue
            if all(s in h for s in subs):
                return i
        return None

    cols = {}
    cols['showroom'] = find('showroom')
    cols['phase'] = find('phase')
    # 'Showrooms' also contains "room", so prefer the distinctive 'Areas' and
    # otherwise look only to the right of the showroom column.
    cols['rooms'] = find('area')
    if cols['rooms'] is None:
        cols['rooms'] = find('room', after=cols['showroom'] if cols['showroom'] is not None else -1)
    cols['removalTeam'] = find('removal', 'team')
    # "Added Manpower / Each Day" appears twice; disambiguate by the team column
    # each one follows.
    cols['removalManpower'] = find('manpower', after=cols['removalTeam'] if cols['removalTeam'] is not None else -1)
    cols['installTeam'] = find('installation', 'team')
    cols['installManpower'] = find('manpower', after=cols['installTeam'] if cols['installTeam'] is not None else -1)
    cols['plannedRemoval'] = find('date', 'removal')
    cols['plannedInstall'] = find('date', 'install')
    cols['plannedCompletion'] = find('date', 'completion')
    cols['initialDuration'] = find('initial', 'duration')
    cols['plannedDuration'] = find('planned', 'duration')
    cols['excludedDates'] = find('excluded')
    cols['agreedTeams'] = find('agreed')
    cols['installationIncharge'] = find('incharge') or find('in charge')

    missing = [k for k in REQUIRED if cols.get(k) is None]
    if missing:
        raise SystemExit(
            f'Could not locate required column(s) {missing} in the header row '
            f'of {ws.title!r}. Headers found: {[h for h in hdr if h]}. '
            f'The sheet layout changed — this needs a human look rather than a '
            f'guess, because reading the wrong column publishes wrong numbers.'
        )
    return cols


def get_token():
    tenant = os.environ['AZURE_TENANT_ID']
    data = {
        'client_id': os.environ['AZURE_CLIENT_ID'],
        'grant_type': 'client_credentials',
        'scope': 'https://graph.microsoft.com/.default',
    }
    secret = os.environ.get('AZURE_CLIENT_SECRET')
    if secret:
        data['client_secret'] = secret
    else:
        # No secret set: exchange the GitHub Actions OIDC token instead.
        gh = requests.get(
            os.environ['ACTIONS_ID_TOKEN_REQUEST_URL'],
            params={'audience': 'api://AzureADTokenExchange'},
            headers={'Authorization':
                     f"Bearer {os.environ['ACTIONS_ID_TOKEN_REQUEST_TOKEN']}"},
            timeout=30,
        )
        gh.raise_for_status()
        data['client_assertion_type'] = \
            'urn:ietf:params:oauth:client-assertion-type:jwt-bearer'
        data['client_assertion'] = gh.json()['value']

    resp = requests.post(
        f'https://login.microsoftonline.com/{tenant}/oauth2/v2.0/token',
        data=data, timeout=30,
    )
    if resp.status_code != 200:
        raise SystemExit(
            f'Token request failed ({resp.status_code}): {resp.text[:500]}\n'
            f'Check AZURE_TENANT_ID / AZURE_CLIENT_ID / AZURE_CLIENT_SECRET.'
        )
    return resp.json()['access_token']


def download_workbook(token):
    url = (f'https://graph.microsoft.com/v1.0/drives/{DRIVE_ID}'
           f'/items/{ITEM_ID}/content')
    resp = requests.get(url, headers={'Authorization': f'Bearer {token}'},
                        timeout=60)
    if resp.status_code != 200:
        raise SystemExit(
            f'Graph file download failed ({resp.status_code}): '
            f'{resp.text[:500]}\n'
            f'The app needs a per-file read grant on this item — see the '
            f'runbook. Check DRIVE_ID/ITEM_ID are still correct.'
        )
    return BytesIO(resp.content)


def _s(v):
    """Text cell -> trimmed str or None. Internal newlines are preserved on
    purpose: excludedDates and agreedTeams are multi-line and the page renders
    them so. Safe because the payload is JSON-encoded."""
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    return s or None


def _num(v):
    if v is None or v == '':
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v) == int(v) else v
    try:
        f = float(str(v).strip())
        return int(f) if f == int(f) else f
    except ValueError:
        return _s(v)


def _date(v):
    """Dates -> ISO yyyy-mm-dd. A literal '-' is a deliberate placeholder the
    page checks for explicitly, so it is passed through unchanged."""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime('%Y-%m-%d')
    return _s(v)


def pick_sheet(wb, want):
    """Exact (whitespace-insensitive) match. Deliberately not a substring
    match: this workbook also holds an older 'Showroom Plan ' sheet that a
    substring search for 'Showroom Plan' would hit first."""
    for nm in wb.sheetnames:
        if nm.strip() == want:
            return wb[nm]
    raise SystemExit(
        f'Sheet {want!r} not found. Sheets present: {wb.sheetnames}. '
        f'The workbook structure may have changed — this needs a human look.'
    )


def parse_plan(ws):
    """Column A names a showroom on the first row of its block; the phase cell
    starts a phase and that row carries the phase-level fields; later rows with
    only a room cell append more rooms to the current phase. All column indexes
    come from header_map, never from fixed positions."""
    cols = header_map(ws)
    width = max(c for c in cols.values() if c is not None) + 1

    def cell(row, key):
        i = cols.get(key)
        return None if i is None else row[i]

    showrooms, cur_sr, cur_ph = [], None, None
    for row in ws.iter_rows(min_row=2, values_only=True):
        row = list(row) + [None] * width
        a_s = _s(cell(row, 'showroom'))
        b_s = _s(cell(row, 'phase'))
        c_s = _s(cell(row, 'rooms'))

        # Require the word "Showroom": that column also carries stray markers
        # such as a bare "TBC" that must not open a new block.
        if a_s and 'showroom' in a_s.lower():
            cur_sr = {'name': ' '.join(a_s.split()), 'phases': []}
            showrooms.append(cur_sr)
            cur_ph = None
        if cur_sr is None:
            continue

        if b_s:
            team = note = None
            d_s = _s(cell(row, 'removalTeam'))
            if d_s:
                if TEAM_RE.match(d_s):
                    team = d_s
                else:
                    note = d_s
            cur_ph = {
                'phase': b_s,
                'rooms': [],
                'removalTeam': team,
                'removalManpower': _num(cell(row, 'removalManpower')),
                'installTeam': _num(cell(row, 'installTeam')),
                'installManpower': _num(cell(row, 'installManpower')),
                'plannedRemoval': _date(cell(row, 'plannedRemoval')),
                'plannedInstall': _date(cell(row, 'plannedInstall')),
                'plannedCompletion': _date(cell(row, 'plannedCompletion')),
                'initialDuration': _s(cell(row, 'initialDuration')),
                'plannedDuration': _s(cell(row, 'plannedDuration')),
                'excludedDates': _s(cell(row, 'excludedDates')),
                'agreedTeams': _s(cell(row, 'agreedTeams')),
                'installationIncharge': _s(cell(row, 'installationIncharge')),
                'note': note,
            }
            cur_sr['phases'].append(cur_ph)

        if c_s:
            if cur_ph is None:
                raise SystemExit(
                    f'Room {c_s!r} appears before any phase in '
                    f'{cur_sr["name"]!r} — sheet layout changed.'
                )
            cur_ph['rooms'].append(c_s)

        # Some phases carry their removal date on a continuation row.
        if not b_s and cur_ph is not None and cur_ph['plannedRemoval'] is None:
            hd = _date(cell(row, 'plannedRemoval'))
            if hd:
                cur_ph['plannedRemoval'] = hd
    return showrooms


def parse_sequence(ws):
    """Finds the header row by its 'Old'/'New' labels and takes the column
    positions from it, so inserted columns cannot shift the values."""
    seq, idx = [], None
    for row in ws.iter_rows(values_only=True):
        vals = [_s(v) for v in row]
        if idx is None:
            low = [(v or '').lower() for v in vals]
            o = next((i for i, v in enumerate(low) if 'old' in v), None)
            n = next((i for i, v in enumerate(low) if 'new' in v), None)
            if o is not None and n is not None:
                r = next((i for i, v in enumerate(low) if 'remark' in v), None)
                idx = (o, n, r)
            continue
        o, n, r = idx
        a = vals[o] if o < len(vals) else None
        b = vals[n] if n < len(vals) else None
        c = vals[r] if r is not None and r < len(vals) else None
        if a and b:
            seq.append({'old': a, 'new': b, 'remark': c})
    if not seq:
        raise SystemExit(
            'Old-New Sequence: could not find an Old/New header row, or it had '
            'no data rows beneath it.'
        )
    return seq


def validate(data):
    srs = data['showrooms']
    names = [s['name'] for s in srs]
    if len(srs) != EXPECTED_SHOWROOMS:
        raise SystemExit(
            f'Expected {EXPECTED_SHOWROOMS} showrooms, parsed '
            f'{len(srs)}: {names}'
        )
    for s in srs:
        if not s['phases']:
            raise SystemExit(f'Showroom {s["name"]!r} parsed with no phases')
        for p in s['phases']:
            if not p['rooms']:
                raise SystemExit(
                    f'{s["name"]!r} phase {p["phase"]!r} parsed with no rooms'
                )
    if not data['sequence']:
        raise SystemExit('Old-New Sequence parsed no rows')
    return names


def main():
    if len(sys.argv) < 3:
        raise SystemExit(__doc__.strip().splitlines()[-1])
    in_path, out_path = sys.argv[1], sys.argv[2]

    wb = openpyxl.load_workbook(download_workbook(get_token()), data_only=True)
    now = dt.datetime.now(dt.timezone(dt.timedelta(hours=4)))  # Dubai / GST, UTC+4, no DST
    data = {
        'showrooms': parse_plan(pick_sheet(wb, PLAN_SHEET)),
        'sequence': parse_sequence(pick_sheet(wb, SEQ_SHEET)),
        'today': now.strftime('%Y-%m-%d'),
        'source': SOURCE_LABEL,
        'generated': now.strftime('%d %b %Y, %H:%M GST'),
    }
    names = validate(data)

    html = open(in_path, encoding='utf-8').read()
    payload = 'const DATA = ' + json.dumps(data, ensure_ascii=False) + ';'
    new_html, n = re.subn(r'const DATA = \{.*?\};', lambda _m: payload,
                          html, count=1, flags=re.DOTALL)
    if n == 0:
        raise SystemExit('DATA block not found/replaced — check the anchor')

    with open(out_path, 'w', encoding='utf-8') as fh:
        fh.write(new_html)

    phases = sum(len(s['phases']) for s in data['showrooms'])
    print(f'OK — wrote {out_path}')
    print(f'Showrooms: {len(names)} ({", ".join(names)})')
    print(f'Phases: {phases} · sequence rows: {len(data["sequence"])}')


if __name__ == '__main__':
    main()
